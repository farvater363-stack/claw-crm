"""One-off import of Рассчеты.xlsx into Claw CRM. Usage: import-excel.py <xlsx> [--dry-run]"""
import datetime
import json
import os
import sys
import urllib.error
import urllib.request

import openpyxl

INSTALLATION_DAYS_AFTER_START = 7

SOURCE = {'OLX': 'OLX', 'Instagram': 'INSTAGRAM', 'Facebook': 'FACEBOOK',
          'Telegram': 'TELEGRAM', 'Через знакомых': 'REFERRAL'}
DISTRICT = {'Алмазарский': 'ALMAZAR', 'Бектемирский': 'BEKTEMIR', 'Мирабадский': 'MIRABAD',
            'Мирзо-Улугбекский': 'MIRZO_ULUGBEK', 'Сергелийский': 'SERGELI', 'Учтепинский': 'UCHTEPA',
            'Чиланзарский': 'CHILANZAR', 'Шайхантахурский': 'SHAYKHANTAKHUR', 'Юнусабадский': 'YUNUSABAD',
            'Яккасарайский': 'YAKKASARAY', 'Яшнабадский': 'YASHNABAD', 'Янгихаётский': 'YANGIHAYOT'}

# (column in «Расходы», grille name, metal)
GRILLE_COLUMNS = [
    ('D', 'Прут', 'ROD'),
    ('G', 'Арматура', 'REBAR'),
    ('H', 'Профиль', 'PROFILE'),
    ('E', 'Мараканд 2', 'PROFILE'),
    ('F', 'Мараканд 3', 'PROFILE'),
    ('I', 'Хайтек', 'PROFILE'),
]
CANOPY_COLUMNS = list('JKLMNOPQRSTU')
REQUEST_TIMEOUT_SECONDS = 60
ENTITY_PLURALS = ['masters', 'designs', 'extraServices', 'orders']


def money(amount):
    return None if amount is None else {'amountMicros': round(amount) * 1_000_000, 'currencyCode': 'UZS'}


def number_or_none(value):
    return value if isinstance(value, (int, float)) else None


def iso_date(value):
    return value.date().isoformat() if isinstance(value, datetime.datetime) else None


def build_payloads(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    data_sheet, costs, orders_sheet = workbook['Данные'], workbook['Расходы'], workbook['Расчеты']

    # Расходы!B2 is the flat price the sheet multiplies m² by.
    flat_price = number_or_none(costs['B2'].value)

    masters = [{'name': cell.value} for (cell,) in data_sheet.iter_rows(min_col=3, max_col=3) if cell.value]

    designs = [
        {
            'name': name,
            'metal': metal,
            'pricePerSquareMeter': money(flat_price),
            'materialCostPerSquareMeter': money(number_or_none(costs[f'{column}2'].value)),
            'manufacturingCostPerSquareMeter': money(number_or_none(costs[f'{column}3'].value)),
            'installationCostPerSquareMeter': money(number_or_none(costs[f'{column}4'].value)),
        }
        for (column, name, metal) in GRILLE_COLUMNS
    ]
    # Grilles the orders mention but «Расходы» has no column for come in without a price.
    priced_names = {design['name'] for design in designs}
    other_names = set()
    for row in range(3, orders_sheet.max_row + 1):
        if orders_sheet[f'O{row}'].value:
            other_names.add(str(orders_sheet[f'O{row}'].value).strip())
    designs += [{'name': name} for name in sorted(other_names - priced_names)]

    extra_services = []
    for column in CANOPY_COLUMNS:
        header = str(costs[f'{column}1'].value)  # e.g. «Коз пласт 50», «Коз туника 60»
        kind = 'пластик' if 'пласт' in header else 'туника'
        width = header.split()[-1]
        cost = sum(number_or_none(costs[f'{column}{row}'].value) or 0 for row in range(2, 6))
        extra_services.append({
            'name': f'Козырёк {kind} {width}',
            'unit': 'PER_RUNNING_METER',
            'kind': 'VISOR',
            'cost': money(cost),
        })

    orders, skipped, warnings = [], [], []
    master_names = {master['name'] for master in masters}
    today = datetime.date.today().isoformat()
    for row in range(3, orders_sheet.max_row + 1):
        cell = lambda column: orders_sheet[f'{column}{row}'].value
        if not isinstance(cell('B'), (int, float)) or not (cell('J') or cell('L')):
            continue
        start = iso_date(cell('E'))
        deadline = (
            (cell('E') + datetime.timedelta(days=INSTALLATION_DAYS_AFTER_START)).date().isoformat()
            if start else None
        )
        total, margin = number_or_none(cell('S')), number_or_none(cell('V'))
        if total is None or margin is None:
            skipped.append(row)
        for column, label, mapping in (('F', 'source', SOURCE), ('G', 'district', DISTRICT)):
            if cell(column) and cell(column) not in mapping:
                warnings.append(f'row {row}: {label} value not mapped, imported empty')
        if cell('AA') and cell('AA') not in master_names:
            warnings.append(f'row {row}: master not in «Данные», imported without a master')
        measurement = iso_date(cell('D'))
        orders.append({
            'number': int(cell('B')),
            'name': f"№{int(cell('B')):04d}",
            'clientName': cell('J'),
            'clientPhone': str(int(cell('K'))) if isinstance(cell('K'), (int, float)) else cell('K'),
            'source': SOURCE.get(cell('F')),
            'district': DISTRICT.get(cell('G')),
            'addressLine': cell('H'),  # the server reserves the field name `address`
            'floor': number_or_none(cell('I')),
            'measurementDate': f'{measurement}T00:00:00+05:00' if measurement else None,
            'productionStartDate': start,
            'installationDeadline': deadline,
            'areaSquareMeters': number_or_none(cell('L')),
            'total': money(total),
            'costTotal': money(total - margin) if total is not None and margin is not None else None,
            'prepayment': money(number_or_none(cell('AB'))),
            'masterName': cell('AA'),
            'status': 'PRODUCTION' if deadline and deadline >= today else 'CLOSED',
        })

    return {'masters': masters, 'designs': designs,
            'extraServices': extra_services, 'orders': orders}, skipped, warnings


def graphql(query, variables):
    request = urllib.request.Request(
        f"{os.environ['TWENTY_API_URL']}/graphql",
        data=json.dumps({'query': query, 'variables': variables}).encode(),
        headers={'Content-Type': 'application/json',
                 'Authorization': f"Bearer {os.environ['TWENTY_API_KEY']}"},
    )
    try:
        with urllib.request.urlopen(request, timeout=REQUEST_TIMEOUT_SECONDS) as response:
            body = json.loads(response.read())
    except urllib.error.HTTPError as error:
        print(error.read().decode(errors='replace'), file=sys.stderr)
        raise
    if body.get('errors'):
        raise RuntimeError(json.dumps(body['errors'], ensure_ascii=False))
    return body['data']


def create_many(mutation_name, input_type, records):
    data = graphql(
        f'mutation($data: [{input_type}!]!) {{ {mutation_name}(data: $data) {{ id name }} }}',
        {'data': records},
    )
    return {record['name']: record['id'] for record in data[mutation_name]}


def abort_if_workspace_not_empty():
    counts = graphql('{ ' + ' '.join(f'{name} {{ totalCount }}' for name in ENTITY_PLURALS) + ' }', {})
    existing = {name: counts[name]['totalCount'] for name in ENTITY_PLURALS if counts[name]['totalCount']}
    if existing:
        sys.exit(f'aborting, the import must run exactly once but the workspace already has {existing}; '
                 'reset or clean these records first')


def main():
    path, dry_run = sys.argv[1], '--dry-run' in sys.argv
    payloads, skipped, warnings = build_payloads(path)

    print({key: len(value) for key, value in payloads.items()})
    print(json.dumps(payloads['orders'][0], ensure_ascii=False, indent=2))
    if skipped:
        print(f'rows without cached total or margin: {skipped}')
    for warning in warnings:
        print(f'warning: {warning}')
    if dry_run:
        return

    abort_if_workspace_not_empty()
    master_ids = create_many('createMasters', 'MasterCreateInput', payloads['masters'])
    create_many('createDesigns', 'DesignCreateInput', payloads['designs'])
    create_many('createExtraServices', 'ExtraServiceCreateInput', payloads['extraServices'])

    orders = []
    for order in payloads['orders']:
        master_name = order.pop('masterName')
        orders.append({**order, 'masterId': master_ids.get(master_name)})
    create_many('createOrders', 'OrderCreateInput', orders)
    print('done')


if __name__ == '__main__':
    main()
